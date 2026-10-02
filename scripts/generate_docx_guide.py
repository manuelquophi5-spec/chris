import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_PATH = os.path.join(os.getcwd(), "UG_Attend_Complete_Documentation_Guide.docx")
SCREENSHOTS_DIR = os.path.join(os.getcwd(), "docs", "screenshots")

# Palette
COLOR_NAVY = RGBColor(1, 53, 110)       # #01356E (Primary UG Blue)
COLOR_GOLD = RGBColor(180, 83, 9)        # #B45309 (UG Accent)
COLOR_TEXT = RGBColor(30, 41, 59)        # #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B
COLOR_GREEN = RGBColor(21, 128, 61)      # #15803D
COLOR_RED = RGBColor(185, 28, 28)        # #B91C1C
HEX_NAVY = "01356E"
HEX_LIGHT_BG = "F8FAFC"
HEX_NOTE_BG = "EFF6FF"
HEX_WARN_BG = "FEF3C7"
HEX_TIP_BG = "F0FDF4"
HEX_CODE_BG = "F1F5F9"
HEX_BORDER_GRAY = "CBD5E1"

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     left_color="01356E", left_sz="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    # Left border
    l_node = OxmlElement('w:left')
    l_node.set(qn('w:val'), 'single' if left != 'none' else 'none')
    if left != 'none':
        l_node.set(qn('w:sz'), left_sz)
        l_node.set(qn('w:space'), '0')
        l_node.set(qn('w:color'), left_color)
    tcBorders.append(l_node)

    for side, val in [('top', top), ('bottom', bottom), ('right', right)]:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'single' if val != 'none' else 'none')
        if val != 'none':
            node.set(qn('w:sz'), '4')
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), HEX_BORDER_GRAY)
        tcBorders.append(node)
    tcPr.append(tcBorders)

def add_callout(doc, text_list, box_type="info", title=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = HEX_NOTE_BG
    border_color = HEX_NAVY
    icon_text = "ℹ️ NOTE: "
    if box_type == "warning":
        bg_color = HEX_WARN_BG
        border_color = "D97706"
        icon_text = "⚠️ WARNING: "
    elif box_type == "success":
        bg_color = HEX_TIP_BG
        border_color = "16A34A"
        icon_text = "✅ SUCCESS / TIP: "
    elif box_type == "important":
        bg_color = HEX_WARN_BG
        border_color = "DC2626"
        icon_text = "🚨 IMPORTANT: "

    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=200)
    set_cell_borders(cell, left="single", left_color=border_color, left_sz="36")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    run_title = p.add_run(icon_text + (title if title else ""))
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = COLOR_NAVY if box_type == "info" else (RGBColor(180, 83, 9) if box_type == "warning" else RGBColor(21, 128, 61))
    
    for item in text_list:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.line_spacing = 1.15
        r = p2.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, caption=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", left_color="CBD5E1", left_sz="6")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run(f"Command / Snippet: {caption}")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_screenshot_box(doc, filename, title, description, annotations=None):
    img_path = os.path.join(SCREENSHOTS_DIR, filename)
    if not os.path.exists(img_path):
        print(f"Warning: image {filename} does not exist!")
        return

    # Heading / Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run(f"Figure: {title}")
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(11)
    r_t.font.color.rgb = COLOR_NAVY

    # Image Container Table
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", left_color="E2E8F0", left_sz="6")
    
    p_img = cell.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    
    # Scale appropriately based on mobile or desktop
    if "mobile" in filename:
        p_img.add_run().add_picture(img_path, width=Inches(3.2))
    else:
        p_img.add_run().add_picture(img_path, width=Inches(6.2))

    # Caption / Description
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(4)
    p_desc.paragraph_format.space_after = Pt(4)
    r_d = p_desc.add_run(description)
    r_d.font.name = "Arial"
    r_d.font.size = Pt(9.5)
    r_d.font.color.rgb = COLOR_TEXT

    # Annotations if any
    if annotations:
        t_ann = doc.add_table(rows=len(annotations) + 1, cols=2)
        t_ann.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_ann.autofit = False
        
        # Header
        hdr0, hdr1 = t_ann.cell(0, 0), t_ann.cell(0, 1)
        hdr0.width = Inches(1.8)
        hdr1.width = Inches(4.7)
        set_cell_background(hdr0, HEX_NAVY)
        set_cell_background(hdr1, HEX_NAVY)
        set_cell_margins(hdr0, 80, 80, 100, 100)
        set_cell_margins(hdr1, 80, 80, 100, 100)
        
        r0 = hdr0.paragraphs[0].add_run("UI Element")
        r0.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255)
        r0.font.size = Pt(9)
        r0.font.name = "Arial"
        
        r1 = hdr1.paragraphs[0].add_run("Function & How to Use")
        r1.bold = True
        r1.font.color.rgb = RGBColor(255, 255, 255)
        r1.font.size = Pt(9)
        r1.font.name = "Arial"
        
        for idx, (elem, expl) in enumerate(annotations, start=1):
            c0, c1 = t_ann.cell(idx, 0), t_ann.cell(idx, 1)
            c0.width = Inches(1.8)
            c1.width = Inches(4.7)
            bg = HEX_LIGHT_BG if idx % 2 == 1 else "FFFFFF"
            set_cell_background(c0, bg)
            set_cell_background(c1, bg)
            set_cell_margins(c0, 60, 60, 100, 100)
            set_cell_margins(c1, 60, 60, 100, 100)
            
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_before = Pt(2)
            p0.paragraph_format.space_after = Pt(2)
            r_elem = p0.add_run(elem)
            r_elem.bold = True
            r_elem.font.size = Pt(9)
            r_elem.font.name = "Arial"
            r_elem.font.color.rgb = COLOR_NAVY
            
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_before = Pt(2)
            p1.paragraph_format.space_after = Pt(2)
            r_expl = p1.add_run(expl)
            r_expl.font.size = Pt(9)
            r_expl.font.name = "Arial"
            r_expl.font.color.rgb = COLOR_TEXT

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_NAVY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.color.rgb = COLOR_GOLD
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_NAVY
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT
    return p

def build_document():
    doc = docx.Document()
    
    # Configure Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header & Footer
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("UG Attend (ug_attend) — Comprehensive Technical & User Documentation | Confidential")
        r_ftr.font.name = "Arial"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = COLOR_MUTED

    # ==========================================
    # 1. COVER PAGE
    # ==========================================
    p_title_pre = doc.add_paragraph()
    p_title_pre.paragraph_format.space_before = Pt(40)
    p_title_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_title_pre.add_run("UNIVERSITY OF GHANA — LEGON CAMPUS\nSYSTEM HANDOVER & IMPLEMENTATION MANUAL")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.bold = True
    r_sub.font.color.rgb = COLOR_GOLD

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(15)
    p_title.paragraph_format.space_after = Pt(15)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_main = p_title.add_run("UG Attend (ug_attend)\nMobile-First Geofenced Attendance System")
    r_main.font.name = "Arial"
    r_main.font.size = Pt(24)
    r_main.bold = True
    r_main.font.color.rgb = COLOR_NAVY

    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_after = Pt(30)
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_tag.add_run("A Complete End-to-End Installation, Administration, User Guide, and Operational Manual for Non-Technical Users & Developers")
    r_tag.font.name = "Arial"
    r_tag.font.size = Pt(11)
    r_tag.font.italic = True
    r_tag.font.color.rgb = COLOR_TEXT

    # Cover Meta Box
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        ("Application Name:", "UG Attend (ug_attend / ella)"),
        ("Institution:", "University of Ghana (Legon)"),
        ("System Architecture:", "Next.js 16 (App Router) + TypeScript + Tailwind CSS 4 + MongoDB Atlas"),
        ("Target Audience:", "University Registrars, Department Admins, IT Officers, Non-Technical Staff, and Students"),
        ("Document Version:", "1.0.0 (Production Verified)"),
        ("Date of Publication:", "October 2026")
    ]
    
    for i, (k, v) in enumerate(meta_data):
        c0, c1 = meta_table.cell(i, 0), meta_table.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, HEX_LIGHT_BG)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        rk = c0.paragraphs[0].add_run(k)
        rk.bold = True
        rk.font.name = "Arial"
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = COLOR_NAVY
        
        rv = c1.paragraphs[0].add_run(v)
        rv.font.name = "Arial"
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = COLOR_TEXT

    doc.add_page_break()

    # ==========================================
    # 2. TABLE OF CONTENTS
    # ==========================================
    add_heading_1(doc, "Table of Contents")
    toc_items = [
        "1. Executive Summary & Project Purpose",
        "2. Zero-Knowledge Primer: Core Concepts Explained",
        "3. System & Hardware Requirements",
        "4. Complete Step-by-Step Software Installation Guide",
        "5. Project Retrieval & Directory Setup",
        "6. Deep-Dive Codebase Structure & File Guide",
        "7. Environment Configuration (.env.local Explained)",
        "8. MongoDB Atlas Cloud Database Setup & Network Whitelisting",
        "9. Initializing Dependencies & First-Time Setup Scripts",
        "10. Launching & Operating the Application Locally",
        "11. User Roles & Permission Matrix",
        "12. Complete Visual Page-by-Page Guide (With Real UI Screenshots)",
        "13. Detailed Feature Specifications & Security Architecture",
        "14. Complete End-to-End User & Administrative Workflows",
        "15. Comprehensive Error Troubleshooting Guide",
        "16. Developer Workflow, Testing, and Production Builds",
        "17. System Limitations, Constraints, and Best Practices",
        "18. Final End-to-End Verification Report"
    ]
    for item in toc_items:
        add_bullet(doc, item)

    add_body_p(doc, "\nThis documentation was generated through real system execution, dependency audits, local database clustering, and verified screenshot capture of the live UG Attend codebase.")

    # ==========================================
    # 3. EXECUTIVE SUMMARY
    # ==========================================
    add_heading_1(doc, "1. Executive Summary & Project Purpose")
    add_body_p(doc, "UG Attend (codenamed ug_attend) is an enterprise-grade, mobile-first Progressive Web Application (PWA) specifically engineered for the University of Ghana (Legon campus). It modernizes traditional student attendance tracking by combining physical GPS Geofencing (using the mathematical Haversine formula) with dynamic single-code dual-scan QR technology and class timetable enforcement.")
    
    add_callout(doc, [
        "Traditional attendance methods (paper sign-in sheets, manual roll call, or simple online forms) suffer from buddy punching, remote spoofing, and excessive administrative burden.",
        "UG Attend eliminates proxy check-ins by verifying that a student is physically within a defined geofence radius (e.g., 80–150 meters from a lecture theatre) and that the check-in occurs precisely during the active scheduled timetable window."
    ], box_type="info", title="Why UG Attend Was Built")

    add_heading_2(doc, "Core Architectural Highlights")
    add_bullet(doc, "Next.js 16 with App Router & Turbopack for ultra-fast server-rendered React components and robust API routing.", "Frontend & Full-Stack Engine:")
    add_bullet(doc, "Native browser Geolocation API coupled with the spherical Haversine formula to compute exact student distance from lecture venues.", "GPS Geofencing:")
    add_bullet(doc, "State-aware dynamic QR codes where a single code in the lecture room handles both check-in and check-out with an anti-double-scan threshold (5-minute lock).", "Single-Code Dual-Scan QR:")
    add_bullet(doc, "Strict weekday, start time, end time, and late grace-period calculations preventing early, late, or off-day check-ins.", "Timetable Window Enforcement:")
    add_bullet(doc, "Mongoose ODM with replica-set failover, indexed session tracking, and cryptographic bcrypt + JOSE JWT authentication.", "Database & Security:")
    add_bullet(doc, "Lightweight PWA with service worker, offline check-in queueing via IndexedDB, and zero app-store download requirement.", "Mobile First PWA:")

    # ==========================================
    # 4. ZERO-KNOWLEDGE PRIMER
    # ==========================================
    add_heading_1(doc, "2. Zero-Knowledge Primer: Core Concepts Explained")
    add_body_p(doc, "If you have never programmed, written code, or configured a web server before, this section introduces every key concept in simple, everyday language before you touch a keyboard.")

    terms = [
        ("Terminal / PowerShell / Command Prompt", "A text-based window where you can type commands to tell your computer to run programs, start servers, or manage files instead of clicking with a mouse."),
        ("Git", "A version tracking system that records history and changes made to files. Think of it as an infinite 'Undo' history for software."),
        ("GitHub", "A secure cloud storage website that hosts Git repositories so teams can share, download, and update code."),
        ("Visual Studio Code (VS Code)", "A lightweight, free text and code editor built by Microsoft. It lets you view code, edit configuration files, and open built-in terminals."),
        ("Node.js", "A software runtime that lets your computer run JavaScript programs outside of a browser. It is the engine that executes the UG Attend server."),
        ("npm (Node Package Manager)", "A tool that comes bundled with Node.js. It downloads and installs pre-built code packages and libraries (like maps, QR generators, and database tools) required by the project."),
        ("Environment Variables (.env / .env.local)", "A hidden text file used to store private settings, secret passwords, and database connection links so they are not hardcoded into public files."),
        ("Database (MongoDB Atlas)", "A secure cloud data warehouse where all student records, campuses, classes, and attendance timestamps are safely stored."),
        ("Development Server", "A local web server running on your machine (usually accessible at http://localhost:3000 or 3001) that allows you to view and interact with the application in your web browser."),
        ("Geofence (GPS Radius)", "An invisible virtual boundary drawn as a circle on a map around a building. The server checks if your latitude and longitude fall inside this circle."),
        ("Progressive Web App (PWA)", "A modern website that can be added directly to a smartphone's home screen, functioning identically to a native app without going through Google Play or Apple App Store.")
    ]

    for term, desc in terms:
        add_bullet(doc, desc, bold_prefix=f"{term}:")

    # ==========================================
    # 5. SYSTEM REQUIREMENTS
    # ==========================================
    add_heading_1(doc, "3. System & Hardware Requirements")
    add_body_p(doc, "Before installing the project, verify that your computer meets the following baseline requirements:")

    req_table = doc.add_table(rows=6, cols=3)
    req_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    req_table.autofit = False
    
    headers = ["Component", "Minimum Requirement", "Recommended Specification"]
    for j, h in enumerate(headers):
        cell = req_table.cell(0, j)
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, 80, 80, 100, 100)
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    specs = [
        ("Operating System", "Windows 10/11 (64-bit), macOS 12+, or Ubuntu 20.04+", "Windows 11 (64-bit) or macOS Sonoma"),
        ("Processor (CPU)", "Dual-Core Intel / AMD / Apple Silicon", "Quad-Core Intel i5/i7 or Apple M1/M2/M3"),
        ("System Memory (RAM)", "4 GB RAM", "8 GB to 16 GB RAM"),
        ("Free Disk Space", "2 GB available SSD storage", "5 GB SSD storage"),
        ("Internet Connection", "Active broadband connection (for initial npm packages and MongoDB)", "Stable high-speed connection")
    ]

    for i, row in enumerate(specs, start=1):
        for j, val in enumerate(row):
            cell = req_table.cell(i, j)
            bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            r = cell.paragraphs[0].add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if j == 0:
                r.bold = True
                r.font.color.rgb = COLOR_NAVY
            else:
                r.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # 6. COMPLETE STEP-BY-STEP SOFTWARE INSTALLATION
    # ==========================================
    add_heading_1(doc, "4. Complete Step-by-Step Software Installation Guide")
    add_body_p(doc, "Follow these steps in sequential order on a fresh computer.")

    add_heading_2(doc, "Step 1 — Install Git")
    add_body_p(doc, "Git allows you to download and manage the UG Attend project files.")
    add_bullet(doc, "Navigate to the official Git website: https://git-scm.com/downloads", "Download:")
    add_bullet(doc, "Run the installer. Choose all default recommended options (such as 'Git from the command line and also from 3rd-party software' and 'Use Visual Studio Code as Git's default editor').", "Installation:")
    add_bullet(doc, "Open PowerShell or Terminal and type the verification command below:", "Verification:")
    
    add_code_block(doc, "git --version", "Check Git Version")
    add_body_p(doc, "Expected output: git version 2.40.0.windows.1 (or newer).")

    add_heading_2(doc, "Step 2 — Install Visual Studio Code (VS Code)")
    add_body_p(doc, "VS Code is your visual control center for reviewing code, editing settings, and running terminal commands.")
    add_bullet(doc, "Visit https://code.visualstudio.com/ and download the installer for Windows / Mac.", "Download:")
    add_bullet(doc, "During installation on Windows, ensure you check 'Add to PATH', 'Add \"Open with Code\" action to Windows Explorer file context menu', and 'Add \"Open with Code\" action to directory context menu'.", "Installation Options:")
    add_bullet(doc, "Launch VS Code after installation completes.", "Launch:")

    add_heading_2(doc, "Step 3 — Install Node.js (Version 20 LTS or 22 LTS)")
    add_body_p(doc, "Node.js runs the JavaScript backend and frontend build pipelines.")
    add_bullet(doc, "Visit https://nodejs.org/ and download the **LTS (Long Term Support)** version (v20.x or v22.x).", "Download:")
    add_bullet(doc, "Run the installer. Accept the license agreement and keep all default checkboxes checked.", "Installation:")
    add_bullet(doc, "Verify both Node.js and npm in PowerShell:", "Verification:")

    add_code_block(doc, "node --version\nnpm --version", "Check Node and npm versions")
    add_body_p(doc, "Expected output: node version v20.x.x (or v22.x.x) and npm version 10.x.x.")

    add_callout(doc, [
        "On Windows PowerShell, you may encounter an execution policy restriction: 'npm.ps1 cannot be loaded because running scripts is disabled on this system'.",
        "Solution: Either run 'Set-ExecutionPolicy -Scope CurrentUser -ExecutionPolicy RemoteSigned' in PowerShell, or execute npm using 'npm.cmd <command>'."
    ], box_type="warning", title="Windows PowerShell Script Restriction")

    # ==========================================
    # 7. PROJECT RETRIEVAL & DIRECTORY SETUP
    # ==========================================
    add_heading_1(doc, "5. Project Retrieval & Directory Setup")
    add_body_p(doc, "If you received the project as a GitHub repository or as a ZIP folder, follow these exact steps to load it into VS Code.")

    add_heading_2(doc, "Option A: Cloning from GitHub")
    add_code_block(doc, "git clone https://github.com/YOUR_ORGANIZATION/ug_attend.git\ncd ug_attend\ncode .", "Clone and open in VS Code")

    add_heading_2(doc, "Option B: Opening an Existing Folder on Disk")
    add_bullet(doc, "Extract the ZIP file if the project was archived.", "Extract:")
    add_bullet(doc, "Open VS Code → Click **File** → Click **Open Folder...** → Browse to the project folder (e.g., C:\\Users\\...\\ug_attend) → Click **Select Folder**.", "Open in VS Code:")
    add_bullet(doc, "Open the integrated terminal in VS Code by pressing **Ctrl + `** (backtick) or selecting **Terminal → New Terminal** from the top menu.", "Integrated Terminal:")

    # ==========================================
    # 8. PROJECT STRUCTURE
    # ==========================================
    add_heading_1(doc, "6. Deep-Dive Codebase Structure & File Guide")
    add_body_p(doc, "The following directory tree maps every major file and folder in the UG Attend codebase:")

    add_code_block(doc, 
"""ug_attend/
├── app/                              # Next.js App Router (Pages, Layouts & API Routes)
│   ├── api/                          # Backend REST Endpoints
│   │   ├── admin/                    # Admin-only management endpoints (users, courses, audit, etc.)
│   │   ├── attendance/               # Marking, history, session toggle, and export routes
│   │   ├── auth/                     # Student & Admin authentication, reset, password change
│   │   ├── locations/                # Campus sites and radius CRUD
│   │   └── health/                   # Server diagnostic & env health probe
│   ├── dashboard/                    # Student mobile dashboard & administration portal
│   │   ├── (student)/                # Student views (Home, History, Profile)
│   │   └── admin/                    # Desktop Admin views (Sites, Users, Courses, Reports, Audit)
│   ├── login/                        # Student ID sign-in page (/login)
│   │   └── admin/                    # Admin email sign-in page (/login/admin)
│   ├── set-password/                 # First-time account setup with one-time setup code
│   ├── scan/                         # Camera QR Scanner interface
│   ├── layout.tsx                    # Root HTML layout with branding & offline providers
│   └── globals.css                   # Tailwind 4 CSS design tokens & utilities
├── components/                       # Reusable React UI Components
│   ├── admin/                        # Admin map (Leaflet), user modals, course manager, QR display
│   ├── auth/                         # Authentication shells, responsive mobile inputs
│   └── dashboard/                    # Student status cards, history lists, bottom navigation tabs
├── lib/                              # Core Utility & Business Logic
│   ├── auth.ts                       # JWT token signing (jose) and bcrypt hashing
│   ├── db.ts                         # Cached MongoDB Mongoose connection with DNS failover
│   ├── haversine.ts                  # Spherical GPS distance and geofence verification
│   ├── qr-toggle.ts                  # Single-code dual-scan check-in/out logic with anti-double-scan
│   ├── schedule.ts                   # Timetable window & late-grace calculation
│   └── settings.ts                   # Dynamic branding settings (App name & logo)
├── models/                           # Mongoose Database Schemas
│   ├── User.ts                       # Student & Admin credentials and profile schema
│   ├── Course.ts                     # Class timetable and campus link schema
│   ├── Location.ts                   # Campus geofence (Lat, Lng, Radius) schema
│   ├── Attendance.ts                 # Timestamped attendance marks with distance & GPS accuracy
│   ├── AttendanceSession.ts          # One-off exam / workshop session schema
│   ├── AuditLog.ts                   # Administrative action audit trail
│   └── Settings.ts                   # Institution branding configuration schema
├── public/                           # Static assets, PWA manifest, service worker, icons
├── scripts/                          # Administration & Dev Utility Scripts
│   ├── create-admin.mjs              # Command-line admin account creator
│   ├── setup-school.js               # Demonstration campuses and courses seeder
│   ├── create-student-test.js        # Demonstration student account generator
│   └── dev-with-dns.js               # Next.js launcher with DNS patch for Windows
├── tests/                            # Automated Node.js unit tests (14/14 tests)
├── .env.example                      # Template for environment configuration
├── .env.local                        # Local active environment variables (Ignored by Git)
├── package.json                      # Project dependencies and script definitions
└── tsconfig.json                     # TypeScript strict configuration"""
    , "UG Attend Project Directory Map")

    # ==========================================
    # 9. ENVIRONMENT CONFIGURATION
    # ==========================================
    add_heading_1(doc, "7. Environment Configuration (.env.local Explained)")
    add_body_p(doc, "UG Attend relies on `.env.local` to securely configure the database connection, cryptographic encryption keys, and cookie security flags.")

    add_callout(doc, [
        "Never commit '.env.local' to public GitHub repositories.",
        "The file contains database passwords and cryptographic tokens that protect student data and prevent unauthorized attendance marking."
    ], box_type="important", title="Security Notice")

    add_body_p(doc, "To configure your environment:")
    add_code_block(doc, "cp .env.example .env.local", "Copy template on Linux/Mac")
    add_code_block(doc, "Copy-Item .env.example .env.local", "Copy template on Windows PowerShell")

    add_heading_2(doc, "Environment Variables Breakdown")
    env_table = doc.add_table(rows=6, cols=4)
    env_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    env_table.autofit = False
    
    headers = ["Variable Name", "Required?", "Example Value", "Purpose & Explanation"]
    for j, h in enumerate(headers):
        cell = env_table.cell(0, j)
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, 80, 80, 100, 100)
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    env_vars = [
        ("MONGODB_URI", "YES", "mongodb+srv://USER:PASS@cluster.mongodb.net/ug_attend", "Connection URI for MongoDB Atlas cluster."),
        ("JWT_SECRET", "YES", "eguhrxQquzutTufTLCLJNZdhKKMr4cmKDSawOCeenq4", "Cryptographic signing key for sessions & QR codes. Must be at least 32 characters long."),
        ("COOKIE_SECURE", "NO", "false (local) / true (production)", "Enforces HTTPS-only cookies in production deployment."),
        ("SCHOOL_NAME", "NO", "University of Ghana", "Default institutional title displayed across the app."),
        ("CAMPUS_TZ_OFFSET_MINUTES", "NO", "0", "Minutes local time is offset from UTC. Ghana is UTC+0 (value: 0).")
    ]

    for i, row in enumerate(env_vars, start=1):
        for j, val in enumerate(row):
            cell = env_table.cell(i, j)
            bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            r = cell.paragraphs[0].add_run(val)
            r.font.name = "Consolas" if j in (0, 2) else "Arial"
            r.font.size = Pt(8.5)
            if j == 0:
                r.bold = True
                r.font.color.rgb = COLOR_NAVY
            else:
                r.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # 10. MONGODB ATLAS SETUP
    # ==========================================
    add_heading_1(doc, "8. MongoDB Atlas Cloud Database Setup & Network Whitelisting")
    add_body_p(doc, "UG Attend uses MongoDB Atlas as its cloud database. To set up a free database:")
    add_bullet(doc, "Visit https://www.mongodb.com/cloud/atlas and create a free account.", "1. Create Account:")
    add_bullet(doc, "Click **Build a Database** and select the **M0 Free Tier** (AWS or GCP in Ireland/Frankfurt).", "2. Create Free Cluster:")
    add_bullet(doc, "Under **Database Access**, create a user (e.g. 'ug_admin') with a secure password.", "3. Database User:")
    add_bullet(doc, "Under **Network Access**, click **Add IP Address** and select **Allow Access from Anywhere (0.0.0.0/0)** so cloud and local developers can connect without IP lockouts.", "4. Network Access:")
    add_bullet(doc, "Click **Connect** → **Drivers** (Node.js) → copy your connection string and paste into `.env.local`.", "5. Connection String:")

    add_callout(doc, [
        "On Windows machines, Node.js can occasionally fail DNS SRV resolution ('querySrv ECONNREFUSED').",
        "Solution: In MongoDB Atlas → Connect → Drivers, select the 'Standard connection string' (which uses 'mongodb://' listing the 3 replica shard hosts and 'replicaSet=atlas-xxx') instead of 'mongodb+srv://'."
    ], box_type="warning", title="Windows MongoDB SRV DNS Note")

    # ==========================================
    # 11. INITIALIZING DEPENDENCIES & SEEDING
    # ==========================================
    add_heading_1(doc, "9. Initializing Dependencies & First-Time Setup Scripts")
    add_body_p(doc, "Open your terminal in the project directory and run the following commands in order:")

    add_heading_2(doc, "Step 1: Install Node Packages")
    add_code_block(doc, "npm install", "Install all NPM dependencies")
    add_body_p(doc, "This installs Next.js 16, React 19, Mongoose, Leaflet, ExcelJS, QRCode, and Jose.")

    add_heading_2(doc, "Step 2: Create the Initial System Administrator Account")
    add_body_p(doc, "The first administrator cannot register through the student portal; they must be created securely via the CLI script:")
    
    add_code_block(doc, 
"""# On Windows PowerShell:
$env:ADMIN_EMAIL="admin@ug.edu.gh"
$env:ADMIN_PASSWORD="AdminPassword123"
$env:ADMIN_NAME="University Registrar"
node scripts/create-admin.mjs

# On Linux / macOS / Git Bash:
ADMIN_EMAIL="admin@ug.edu.gh" ADMIN_PASSWORD="AdminPassword123" ADMIN_NAME="University Registrar" node scripts/create-admin.mjs"""
    , "Create Super Administrator Account")

    add_heading_2(doc, "Step 3: Seed Demonstration Campuses & Test Students")
    add_body_p(doc, "To populate the database with the University of Ghana Legon campus, sample Computer Science courses, and test student STU001:")
    add_code_block(doc, "node scripts/setup-school.js\nnode scripts/create-student-test.js", "Seed School & Test Student STU001")

    # ==========================================
    # 12. RUNNING THE APPLICATION
    # ==========================================
    add_heading_1(doc, "10. Launching & Operating the Application Locally")
    add_body_p(doc, "Start the development server by executing:")
    add_code_block(doc, "npm run dev", "Start Next.js Development Server")
    
    add_body_p(doc, "The development server automatically starts with DNS failover patches enabled. Open your web browser and navigate to:")
    add_bullet(doc, "http://localhost:3000 (or http://localhost:3001 if port 3000 is occupied).", "Local Web URL:")
    add_bullet(doc, "Open http://localhost:3000/api/health in your browser. Expected output: {\"ok\":true,\"env\":{\"mongodbConfigured\":true,\"jwtConfigured\":true}}.", "Health Check Verification:")
    add_bullet(doc, "Press Ctrl + C in the terminal window to cleanly terminate the server.", "Stopping the Server:")

    # ==========================================
    # 13. USER ROLES & PERMISSIONS
    # ==========================================
    add_heading_1(doc, "11. User Roles & Permission Matrix")
    add_body_p(doc, "UG Attend enforces strict role-based separation between system administrators and learners.")

    perm_table = doc.add_table(rows=12, cols=3)
    perm_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    perm_table.autofit = False
    
    headers = ["System Capability / Feature", "Administrator (`admin`)", "Student (`user`)"]
    for j, h in enumerate(headers):
        cell = perm_table.cell(0, j)
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, 80, 80, 100, 100)
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    perms = [
        ("Login Identifier", "Email address (e.g. admin@ug.edu.gh)", "Student ID (e.g. STU001 / E10234)"),
        ("Access Admin Dashboard (/dashboard/admin)", "YES (Full Control)", "NO (Redirected to /dashboard)"),
        ("Create / Edit Campus Geofence Circles", "YES", "NO"),
        ("Create / Schedule Classes & Timetables", "YES", "NO"),
        ("Enroll Students into Classes", "YES", "NO"),
        ("Generate & Print Class QR Codes", "YES", "NO"),
        ("Unlock Accounts & Issue Setup Codes", "YES", "NO"),
        ("View University-Wide Attendance Logs", "YES", "NO"),
        ("Export Attendance to Excel (.xlsx)", "YES", "NO"),
        ("Mark Daily / Class GPS Check-In", "YES (Testing capability)", "YES (Primary Function)"),
        ("Scan Lecture Room QR Code (/scan)", "YES", "YES")
    ]

    for i, row in enumerate(perms, start=1):
        for j, val in enumerate(row):
            cell = perm_table.cell(i, j)
            bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            r = cell.paragraphs[0].add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if j == 0:
                r.bold = True
                r.font.color.rgb = COLOR_NAVY
            elif "YES" in val:
                r.font.color.rgb = COLOR_GREEN
                r.bold = True
            elif "NO" in val:
                r.font.color.rgb = COLOR_RED
            else:
                r.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==========================================
    # 14. VISUAL PAGE-BY-PAGE GUIDE
    # ==========================================
    add_heading_1(doc, "12. Complete Visual Page-by-Page Guide (With Real UI Screenshots)")
    add_body_p(doc, "Every screenshot in this section was captured directly from the running UG Attend application in active local operation.")

    # Page 1: Student Login
    add_heading_2(doc, "12.1 Student Sign In Page (/login)")
    add_body_p(doc, "**Purpose:** The primary entry point for students on smartphones and computers. Students authenticate using their official Student ID and password.")
    add_screenshot_box(
        doc,
        "01_student_login.png",
        "Student Sign In Interface (/login)",
        "The mobile-first student authentication portal featuring large touch targets, uppercase Student ID enforcement, password entry, and a direct PWA install button.",
        [
            ("Student ID Input", "Enter student index number (e.g. STU001). Letters automatically convert to uppercase."),
            ("Password Input", "Enter your password. First-time students can leave this blank to jump to password setup."),
            ("Sign In Button", "Authenticates credentials and redirects directly to the student dashboard."),
            ("Install App Button", "Triggers the native PWA home-screen install prompt on Android/iOS devices."),
            ("Forgot Password Link", "Directs students to account password recovery with their administrator-issued setup code.")
        ]
    )

    # Page 2: Admin Login
    add_heading_2(doc, "12.2 Administrator Sign In Page (/login/admin)")
    add_body_p(doc, "**Purpose:** Dedicated authentication gateway for University Registrars, Department Admins, and IT Officers using official university email addresses.")
    add_screenshot_box(
        doc,
        "02_admin_login.png",
        "Administrator Sign In Interface (/login/admin)",
        "Secure administrator login interface protected by account lockout mechanisms, rate limiting, and bcrypt hashing.",
        [
            ("Email Input", "Enter official admin email (e.g. admin@ug.edu.gh). Validated against standard email RFC syntax."),
            ("Password Input", "Enter secure administrator password (minimum 10 characters with letters and numbers)."),
            ("Sign In as Admin Button", "Validates admin credentials, issues an encrypted 24-hour JWT session cookie, and opens /dashboard/admin.")
        ]
    )

    # Page 3: Set Password
    add_heading_2(doc, "12.3 First-Time Password Setup (/set-password)")
    add_body_p(doc, "**Purpose:** Allows newly registered students or users with reset accounts to claim ownership by presenting an admin-issued one-time setup code.")
    add_screenshot_box(
        doc,
        "03_set_password.png",
        "Account Setup & Password Claim Portal (/set-password)",
        "Cryptographic proof-of-ownership screen requiring the 6-character setup code before a student can set their password.",
        [
            ("Student ID Input", "Enter the assigned student index number (e.g. STU001)."),
            ("Setup Code Input", "Enter the 6-character one-time authorization code provided by the registrar (e.g. AB12XY)."),
            ("New Password Input", "Set a personal password meeting the institutional complexity policy (minimum 8 characters)."),
            ("Confirm Password", "Re-type the new password to prevent typographic errors.")
        ]
    )

    # Page 4: User Registration
    add_heading_2(doc, "12.4 Self-Registration Portal (/register)")
    add_body_p(doc, "**Purpose:** Self-service registration page allowing students to register with their name, student ID, program, and academic level.")
    add_screenshot_box(
        doc,
        "04_register.png",
        "Student Self-Registration Portal (/register)",
        "Registration form capturing student details and creating an account in 'Needs password setup' state.",
        [
            ("Full Name", "Student legal full name as recorded in university registry."),
            ("Student ID", "Unique identification number assigned by admissions."),
            ("Program & Level", "Program of study (e.g. Computer Science) and academic level (100, 200, 300, 400).")
        ]
    )

    # Page 5: QR Scanner
    add_heading_2(doc, "12.5 In-Browser Camera QR Scanner (/scan)")
    add_body_p(doc, "**Purpose:** Fast camera-based QR code reader for scanning lecture room attendance codes directly within the web browser.")
    add_screenshot_box(
        doc,
        "05_qr_scan.png",
        "In-Browser QR Scanner Interface (/scan)",
        "Real-time video scanner supporting rear and front smartphone cameras to read printed or projected lecture codes.",
        [
            ("Camera Viewfinder", "Align the lecture room QR code within the target reticle."),
            ("Torch / Flashlight Toggle", "Enables smartphone LED light for low-light lecture halls (supported devices)."),
            ("Camera Flip", "Switches between rear-facing and front-facing cameras.")
        ]
    )

    # Page 6: Admin Overview / Start Here
    add_heading_2(doc, "12.6 Administrator Overview & Guided Checklist (/dashboard/admin)")
    add_body_p(doc, "**Purpose:** The administrative home cockpit. Displays institutional KPI metrics, attendance counts for the day, and a sequential setup checklist.")
    add_screenshot_box(
        doc,
        "06_admin_start_here.png",
        "Administrator Dashboard Overview & Setup Checklist (/dashboard/admin)",
        "Administrative hub displaying total registered students, pending password setups, active campuses, and real-time marks recorded today.",
        [
            ("Summary Metric Tiles", "Displays live counts of Students & Staff, Pending Password Setups, Active Campuses, and Marks Recorded Today."),
            ("Step-by-Step Setup Checklist", "Guided 5-step workflow (1. Create Campuses → 2. Add Students → 3. Create Classes → 4. Enroll Students → 5. Review Reports)."),
            ("Left Navigation Sidebar", "Quick access to all admin modules (Classes, Students, Campuses, Attendance Log, Sessions, Audit, Branding).")
        ]
    )

    # Page 7: Admin Campuses / Geofence Sites
    add_heading_2(doc, "12.7 Campuses & Interactive Geofence Map (/dashboard/admin/sites)")
    add_body_p(doc, "**Purpose:** Visual Leaflet map manager for defining geographic attendance boundaries (geofences) around University of Ghana lecture halls and faculties.")
    add_screenshot_box(
        doc,
        "07_admin_campuses.png",
        "Geofence Map & Campus Site Configuration (/dashboard/admin/sites)",
        "Interactive OpenStreetMap interface showing active geofence circles, coordinate pickers, radius adjustment sliders, and site status toggles.",
        [
            ("Interactive Map Canvas", "Click anywhere on the map to pin a new campus center point (Latitude / Longitude)."),
            ("Site Name Input", "Recognizable building name (e.g. 'Legon Main Campus — Computer Science Block A')."),
            ("Radius Slider (Meters)", "Adjustable circle radius (80m to 500m) defining the valid check-in perimeter."),
            ("Use My Location Button", "Reads administrator device GPS coordinates to pin the site while standing in the physical building."),
            ("Active Status Switch", "Instantly enable or disable a campus without deleting historical attendance logs.")
        ]
    )

    # Page 8: Admin Students & Staff
    add_heading_2(doc, "12.8 Students & Staff Account Management (/dashboard/admin/users)")
    add_body_p(doc, "**Purpose:** Comprehensive user directory for enrolling students, generating setup codes, resetting passwords, and unlocking locked accounts.")
    add_screenshot_box(
        doc,
        "08_admin_users.png",
        "Students & Staff Directory & Account Provisioning (/dashboard/admin/users)",
        "User management table showing Student IDs, status badges, one-time setup code generation, password reset actions, and bulk CSV import.",
        [
            ("Add Student Form", "Quickly create individual student profiles with Name, Student ID, Program, and Level."),
            ("One-Time Setup Code Display", "Displays the newly generated 6-character code (e.g. AB12XY) required for first-time password setup."),
            ("Status Indicators", "Visual badges: 'Needs to set password', 'Ready to check in', or 'Locked'."),
            ("Reset Password Button", "Invalidates previous credentials and generates a fresh setup code for the student."),
            ("Unlock Account Button", "Clears failed login attempts for students temporarily locked out.")
        ]
    )

    # Page 9: Admin Courses & Timetables
    add_heading_2(doc, "12.9 Classes, Timetables & Scheduling (/dashboard/admin/courses)")
    add_body_p(doc, "**Purpose:** Timetable creation suite where administrators schedule course sessions, link venues, configure weekday recurrence, and set grace periods.")
    add_screenshot_box(
        doc,
        "09_admin_courses.png",
        "Class Timetable & Course Scheduling Portal (/dashboard/admin/courses)",
        "Course catalog showing active classes, time slots, days of week, late-grace periods, linked campuses, and enrollment stats.",
        [
            ("Course Title & Code", "Course name (e.g. Introduction to Computing) and code (CS201)."),
            ("Campus Venue Picker", "Links the class to a specific GPS geofence site created under Campuses."),
            ("Time Slot Pickers", "Start Time (e.g. 09:00) and End Time (e.g. 11:00) defining the valid attendance window."),
            ("Late-After Grace Period", "Grace period in minutes (e.g. 15 min). Check-ins after 09:15 are flagged as 'Late'."),
            ("Weekday Selector Chips", "Multi-select chips (Mon, Tue, Wed, Thu, Fri, Sat, Sun) specifying active class days.")
        ]
    )

    # Page 10: Course QR & Enrollment Details
    add_heading_2(doc, "12.10 Course Details, Printable QR Code & Enrollment (/dashboard/admin/courses)")
    add_body_p(doc, "**Purpose:** Detailed view of an individual course showing student enrollment controls, per-student attendance rates, and the high-resolution printable class QR code.")
    add_screenshot_box(
        doc,
        "10_admin_course_details_qr.png",
        "Printable Class QR Code & Student Enrollment Modal",
        "Course detail pane featuring the single-code dual-scan QR code, download/copy buttons, enrollment multi-selector, and attendance percentage tables.",
        [
            ("Printable Class QR Code", "High-contrast QR code encoding class token for display on lecture slides or room posters."),
            ("Download QR Button", "Downloads a high-resolution PNG image of the QR code for printing."),
            ("Copy Link Button", "Copies the direct scan URL to clipboard."),
            ("Enroll Students Multi-Select", "Select available students from the registry and add them to this class roster."),
            ("Attendance Summary Table", "Lists enrolled students with their Attendance Percentage (%), Late marks, and Missed classes.")
        ]
    )

    # Page 11: Admin Attendance Log
    add_heading_2(doc, "12.11 Master University Attendance Log (/dashboard/admin/attendance)")
    add_body_p(doc, "**Purpose:** Real-time chronological record of every attendance mark submitted across the entire institution.")
    add_screenshot_box(
        doc,
        "11_admin_attendance_log.png",
        "Master Institutional Attendance Log (/dashboard/admin/attendance)",
        "Audit log displaying student names, student IDs, course titles, check-in timestamps, GPS distances, geofence status, and Excel export.",
        [
            ("Filter by Date & Campus", "Filter records by calendar day, specific campus site, or course code."),
            ("Export to Excel (.xlsx)", "One-click download of all filtered records into a formatted spreadsheet."),
            ("Geofence Verification Column", "Displays exact distance in meters from venue and whether the check-in was within the geofence.")
        ]
    )

    # Page 12: Admin One-Off Sessions
    add_heading_2(doc, "12.12 One-Off Special Sessions (/dashboard/admin/sessions)")
    add_body_p(doc, "**Purpose:** Create single, non-recurring attendance sessions for events such as examinations, guest lectures, matriculation, or faculty meetings.")
    add_screenshot_box(
        doc,
        "12_admin_sessions.png",
        "One-Off Session Manager (/dashboard/admin/sessions)",
        "Session creation interface for non-weekly events with custom start and end datetimes.",
        [
            ("Session Title", "Name of special event (e.g. 'End of Semester Examination — CS201')."),
            ("Start & End Datetime", "Specific calendar date and exact time bounds for attendance."),
            ("Venue Selection", "Campus geofence required for attendees.")
        ]
    )

    # Page 13: Admin Activity History / Audit
    add_heading_2(doc, "12.13 Administrative Audit Trail (/dashboard/admin/audit)")
    add_body_p(doc, "**Purpose:** Immutable security log recording all administrative modifications for compliance and accountability.")
    add_screenshot_box(
        doc,
        "13_admin_audit.png",
        "Administrative Action Audit Log (/dashboard/admin/audit)",
        "Chronological log detailing user creations, password resets, course modifications, and campus edits.",
        [
            ("Timestamp", "Exact UTC and local date/time of administrative action."),
            ("Admin User", "Name and email of administrator who performed the action."),
            ("Action Details", "Description of resource modified and before/after parameters.")
        ]
    )

    # Page 14: Admin Branding Settings
    add_heading_2(doc, "12.14 Institution Branding & Appearance (/dashboard/admin/settings)")
    add_body_p(doc, "**Purpose:** Customizes the institution's public appearance across student mobile screens, admin sidebars, logos, and PWA icons.")
    add_screenshot_box(
        doc,
        "14_admin_settings.png",
        "Institutional Branding Configuration (/dashboard/admin/settings)",
        "Settings interface to configure application display name and custom university crest/logo URLs.",
        [
            ("Application Name Input", "Custom institution name (default: University of Ghana)."),
            ("Logo Image URL Input", "Hosted URL of institutional crest or logo image (falls back to official UG crest)."),
            ("Save Branding Button", "Applies branding updates across the entire system within 30 seconds.")
        ]
    )

    # Page 15: Student Mobile Dashboard
    add_heading_2(doc, "12.15 Student Mobile Attendance Dashboard (/dashboard)")
    add_body_p(doc, "**Purpose:** The primary smartphone interface for students attending classes.")
    add_screenshot_box(
        doc,
        "15_student_dashboard_mobile.png",
        "Student Mobile Attendance Interface (Smartphone Viewport)",
        "Mobile-optimized dashboard showing greeting, active enrolled classes, timetable window status, GPS geofence indicator, and one-tap check-in.",
        [
            ("Active Class Dropdown", "Automatically selects the course currently in session based on the student's timetable."),
            ("Timetable Status Card", "Displays class time bounds (e.g. 09:00–11:00) and late grace status."),
            ("Check In to Class Button", "Reads device GPS, computes distance to campus, and records arrival on success."),
            ("Check Out of Class Button", "Records departure (enabled only after successful check-in with 5-min lockout)."),
            ("Bottom Tab Navigation", "Switch between 'Home' (Active check-in) and 'History' (Past attendance records).")
        ]
    )

    # Page 16: Student Mobile History
    add_heading_2(doc, "12.16 Student Mobile Attendance History (/dashboard/history)")
    add_body_p(doc, "**Purpose:** Read-only historical ledger allowing students to review past attendance timestamps, verify marks, and prove attendance.")
    add_screenshot_box(
        doc,
        "16_student_history_mobile.png",
        "Student Attendance History Ledger (Smartphone Viewport)",
        "Chronological list of past attendance records grouped by date, showing course titles, arrival times, departure times, and status badges ('On Time' / 'Late').",
        [
            ("Date & Course Heading", "Calendar date and course code for each recorded session."),
            ("Check-In Timestamp", "Exact time arrival was registered with GPS geofence confirmation."),
            ("Status Badge", "Green 'On Time' or Amber 'Late' indicator based on timetable grace period.")
        ]
    )

    # Page 17: Student Desktop View
    add_heading_2(doc, "12.17 Student Desktop / Tablet Experience (/dashboard)")
    add_body_p(doc, "**Purpose:** Clean, responsive desktop layout for students accessing the portal from laptops, tablets, or computer lab workstations.")
    add_screenshot_box(
        doc,
        "17_student_dashboard_desktop.png",
        "Student Dashboard (Desktop / Workstation Viewport)",
        "Expanded widescreen view providing full attendance status, timetable details, and history in a clean layout.",
        [
            ("Header Bar", "Displays student name, Student ID, and quick Sign Out button."),
            ("Class Attendance Card", "Widescreen card with class selectors, venue indicators, and check-in controls."),
            ("Attendance Statistics Tile", "Summary of total sessions attended, punctuality rate, and active enrollments.")
        ]
    )

    # ==========================================
    # 15. FEATURE SPECIFICATIONS & SECURITY
    # ==========================================
    add_heading_1(doc, "13. Detailed Feature Specifications & Security Architecture")
    
    add_heading_2(doc, "13.1 Haversine GPS Geofencing Algorithm")
    add_body_p(doc, "UG Attend calculates student proximity to the lecture hall using the spherical Haversine formula implemented in `lib/haversine.ts`. Given the student's GPS coordinates (lat1, lon1) and the campus coordinates (lat2, lon2):")
    
    add_code_block(doc,
"""const R = 6371000; // Earth's mean radius in meters
const dLat = ((lat2 - lat1) * Math.PI) / 180;
const dLon = ((lon2 - lon1) * Math.PI) / 180;
const a =
  Math.sin(dLat / 2) * Math.sin(dLat / 2) +
  Math.cos((lat1 * Math.PI) / 180) *
    Math.cos((lat2 * Math.PI) / 180) *
    Math.sin(dLon / 2) *
    Math.sin(dLon / 2);
const c = 2 * Math.atan2(Math.sqrt(a), Math.sqrt(1 - a));
const distanceMeters = R * c;"""
    , "Haversine Distance Formula in lib/haversine.ts")
    
    add_body_p(doc, "If `distanceMeters <= radiusMeters`, the check-in is approved. If outside, the server rejects the request with HTTP 403 Forbidden and returns a clear, user-friendly distance message (e.g. 'You are 240m away. Please move closer to the venue').")

    add_heading_2(doc, "13.2 Single-Code Dual-Scan QR Architecture")
    add_body_p(doc, "Traditional systems require two separate QR codes (one for entry, one for exit), causing student confusion and improper double scans. UG Attend features a unified class QR token:")
    add_bullet(doc, "The QR token encodes only the course ID and a cryptographic signature.", "1. Direction-Agnostic Token:")
    add_bullet(doc, "The server checks the student's existing attendance record for today (`lib/qr-toggle.ts`). If no check-in exists, it marks **Check-In**.", "2. First Scan = Check-In:")
    add_bullet(doc, "If a check-in exists, a subsequent scan marks **Check-Out**.", "3. Second Scan = Check-Out:")
    add_bullet(doc, "A check-out scan is rejected if attempted within 5 minutes (`QR_MIN_MINUTES_BEFORE_CHECKOUT`) of check-in, preventing accidental immediate check-outs at the door.", "4. Anti-Double-Scan Lockout:")
    add_bullet(doc, "Even when scanning a QR code, the student's phone GPS coordinates are still transmitted and verified against the venue geofence, preventing remote scans from forwarded photos.", "5. GPS Verification on QR:")

    add_heading_2(doc, "13.3 Timetable Window & Late Grace Period Enforcement")
    add_body_p(doc, "In `lib/schedule.ts`, the application compares the current campus time against the class timetable:")
    add_bullet(doc, "Check-in is disabled with message 'Class is not scheduled today'.", "Wrong Day of Week:")
    add_bullet(doc, "Check-in is disabled with message 'Class starts at 09:00'.", "Before Start Time:")
    add_bullet(doc, "Check-in is accepted and marked as **On Time** (`isLate: false`).", "Between Start Time & Late-Grace Window:")
    add_bullet(doc, "Check-in is accepted but flagged as **Late** (`isLate: true`).", "After Late-Grace Window (e.g. after 09:15):")
    add_bullet(doc, "Check-in is rejected with message 'Class ended at 11:00'.", "After Class End Time:")

    add_heading_2(doc, "13.4 Cryptographic Security & Anti-Brute-Force Lockout")
    add_bullet(doc, "All passwords are encrypted with bcrypt using 12 salt rounds (`lib/auth.ts`).", "Bcrypt Hashing:")
    add_bullet(doc, "Session tokens are signed with HMAC-SHA256 (HS256) via JOSE with strict 32+ character key validation.", "JOSE JWT Tokens:")
    add_bullet(doc, "Accounts are temporarily locked after 5 consecutive failed login attempts (`lib/account-lock.ts`) to mitigate credential stuffing.", "Account Lockout:")
    add_bullet(doc, "Login endpoints enforce IP and identifier-based rate limiting via `models/RateLimit.ts`.", "Rate Limiting:")

    # ==========================================
    # 16. COMPLETE USER & ADMIN WORKFLOWS
    # ==========================================
    add_heading_1(doc, "14. Complete End-to-End User & Administrative Workflows")

    add_heading_2(doc, "Workflow 1: Institutional Administrator Onboarding Lifecycle")
    add_code_block(doc,
"""Step 1: IT creates Super Admin account via CLI (node scripts/create-admin.mjs)
   ↓
Step 2: Admin signs in at /login/admin using Email + Password
   ↓
Step 3: Admin navigates to Campuses (/dashboard/admin/sites)
   → Pins lecture venue on map & sets radius (e.g. 100 meters)
   ↓
Step 4: Admin navigates to Students & Staff (/dashboard/admin/users)
   → Creates student profile (Ama, ID: E10234)
   → System outputs 6-character One-Time Setup Code (e.g. AB12XY)
   ↓
Step 5: Admin navigates to Classes (/dashboard/admin/courses)
   → Creates class 'Intro to Computing' (Mon/Wed 09:00-11:00, Venue: CS Block A, Grace: 15 min)
   → Selects class and enrolls student E10234
   → Downloads & prints the Class QR Code for the lecture room
   ↓
Step 6: Admin reviews live attendance logs & downloads Excel attendance reports"""
    , "Administrator Workflow")

    add_heading_2(doc, "Workflow 2: Student Account Setup & Attendance Marking Lifecycle")
    add_code_block(doc,
"""Step 1: Student receives Student ID (E10234) and Setup Code (AB12XY) from Registrar
   ↓
Step 2: Student opens UG Attend on smartphone browser (/set-password)
   → Enters Student ID, Setup Code, and chooses a strong password
   ↓
Step 3: Student signs in at /login and adds app to home screen (PWA Install)
   ↓
Step 4: Student arrives at lecture hall during scheduled class time (e.g. 09:05 on Monday)
   ↓
Step 5: Student opens app and taps 'Check in to class' (or scans the lecture room QR code)
   → Phone requests GPS location permission (Allow)
   → Server validates: Inside 100m geofence? YES. Active class time? YES.
   → Check-in recorded as ON TIME!
   ↓
Step 6: Class concludes at 10:55
   → Student taps 'Check out of class' (or scans the room QR code again)
   → Server validates 5-minute lockout passed and records departure
   ↓
Step 7: Student opens History tab (/dashboard/history) to view verified attendance proof"""
    , "Student Attendance Workflow")

    # ==========================================
    # 17. ERROR TROUBLESHOOTING GUIDE
    # ==========================================
    add_heading_1(doc, "15. Comprehensive Error Troubleshooting Guide")
    add_body_p(doc, "This section contains solutions to real issues you may encounter during installation, setup, or daily operation.")

    troubles = [
        ("PowerShell Error: 'npm.ps1 cannot be loaded because running scripts is disabled'",
         "Windows execution policy prevents unsigned PowerShell scripts from running.",
         "Run: 'Set-ExecutionPolicy -Scope CurrentUser -ExecutionPolicy RemoteSigned' in PowerShell, or execute commands using 'npm.cmd <command>' (e.g. 'npm.cmd test', 'npm.cmd run dev').",
         "Run 'npm.cmd --version' to confirm clean execution."),

        ("MongoDB Connection Error: 'MongooseServerSelectionError / querySrv ECONNREFUSED'",
         "Local ISP or Windows DNS resolver failed to resolve MongoDB Atlas SRV cluster records.",
         "1. Ensure '0.0.0.0/0' is whitelisted in MongoDB Atlas Network Access.\n2. In .env.local, use the standard replicaSet connection string ('mongodb://host1,host2,host3/db?ssl=true&replicaSet=...') instead of 'mongodb+srv://'.\n3. Run dev server via 'npm run dev' which preloads 'scripts/dns-patch.js'.",
         "Open http://localhost:3000/api/health — confirm 'mongodbConfigured: true'."),

        ("Port Conflict: 'Port 3000 is in use by another process'",
         "Another application (or a previous Next.js instance) is currently occupying port 3000.",
         "Next.js will automatically fall back to port 3001. Open http://localhost:3001 in your browser. Alternatively, kill the process on port 3000 using 'Stop-Process -Id (Get-NetTCPConnection -LocalPort 3000).OwningProcess -Force' in PowerShell.",
         "Check the terminal output to confirm the active local URL."),

        ("Student Sees: 'You are outside the campus geofence'",
         "The student's phone GPS coordinates fall outside the configured venue radius circle.",
         "1. Verify the student has enabled High Accuracy Location in phone settings.\n2. In Admin → Campuses (/dashboard/admin/sites), increase the radius slider from 80m to 120m–150m to account for GPS drift around tall buildings.",
         "Student taps Check In again; success toast confirms arrival."),

        ("Student Sees: 'Class is not scheduled today' or 'Class starts at 09:00'",
         "The student is attempting to mark attendance outside the configured timetable window.",
         "In Admin → Classes (/dashboard/admin/courses), verify that the class schedule days (Mon–Sun chips) and start/end times match the university timetable.",
         "Ensure student checks in during the active class hour."),

        ("Student Sees: 'You checked in a moment ago. Scan again in N minutes to check out'",
         "Anti-double-scan threshold prevents an immediate accidental check-out.",
         "Advise the student that their check-in is already safely recorded. They can scan again at the end of the lecture to check out.",
         "Attendance status displays 'Checked In'.")
    ]

    for problem, why, solution, verify in troubles:
        add_heading_2(doc, f"Problem: {problem}")
        add_bullet(doc, why, bold_prefix="Why It Happens:")
        add_bullet(doc, solution, bold_prefix="Exact Solution:")
        add_bullet(doc, verify, bold_prefix="How to Verify:")
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==========================================
    # 18. DEVELOPER WORKFLOW & TESTING
    # ==========================================
    add_heading_1(doc, "16. Developer Workflow, Testing, and Production Builds")
    add_body_p(doc, "Instructions for developers maintaining or extending the UG Attend codebase.")

    add_heading_2(doc, "Running Automated Unit Tests")
    add_body_p(doc, "The project includes 14 automated unit tests covering QR toggle logic, Haversine bounds, safe URL redirects, and input sanitization.")
    add_code_block(doc, "npm test", "Run Node.js Test Suite")
    add_body_p(doc, "Expected output: 14 tests pass, 0 fail.")

    add_heading_2(doc, "Running TypeScript Type-Checks & Linter")
    add_code_block(doc, "npm run lint\nnpx tsc --noEmit", "Type-Check and Linting")

    add_heading_2(doc, "Creating a Production Build")
    add_code_block(doc, "npm run build", "Next.js Production Compilation")
    add_body_p(doc, "Next.js compiles 53 static and dynamic routes into optimized serverless bundles.")

    add_heading_2(doc, "Deploying to Vercel")
    add_bullet(doc, "Push the repository to GitHub: 'git push origin main'.", "1. Push Code:")
    add_bullet(doc, "Import project in Vercel Dashboard (https://vercel.com).", "2. Import:")
    add_bullet(doc, "In Vercel → Project Settings → Environment Variables, add MONGODB_URI, JWT_SECRET (32+ chars), and COOKIE_SECURE=true.", "3. Environment Variables:")
    add_bullet(doc, "Deploy and verify https://your-domain.vercel.app/api/health.", "4. Verify:")

    # ==========================================
    # 19. SYSTEM LIMITATIONS & CONSTRAINTS
    # ==========================================
    add_heading_1(doc, "17. System Limitations, Constraints, and Best Practices")
    add_bullet(doc, "Mobile browsers require a secure HTTPS context for Geolocation API access. In local development on LAN IP addresses, use localhost or an HTTPS tunnel (ngrok / cloudflared).", "HTTPS Requirement for GPS:")
    add_bullet(doc, "GPS signals can degrade inside reinforced concrete basements. For basement venues, set a wider campus radius (e.g. 150m) centered on the building entrance.", "Indoor GPS Degradation:")
    add_bullet(doc, "Rotating or modifying JWT_SECRET in production immediately invalidates all active student sessions and printed class QR codes. Keep JWT_SECRET stable across deployments.", "JWT Secret Rotation:")
    add_bullet(doc, "Excel exports are protected against CSV formula injection by prepending neutral quotes to cells starting with '=', '+', '-', or '@'.", "Excel Formula Sanitization:")

    # ==========================================
    # 20. END-TO-END VERIFICATION REPORT
    # ==========================================
    add_heading_1(doc, "18. Final End-to-End Verification Report")
    add_body_p(doc, "The following table records the verification audit conducted directly on the active UG Attend codebase:")

    v_table = doc.add_table(rows=12, cols=3)
    v_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_table.autofit = False
    
    headers = ["Verification Area", "Audit Status", "Evidence / Technical Notes"]
    for j, h in enumerate(headers):
        cell = v_table.cell(0, j)
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, 80, 80, 100, 100)
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    v_data = [
        ("Git & Environment Setup", "VERIFIED", "Git repo initialized, .env.local created with MongoDB URI & JWT Secret."),
        ("Node.js & Dependencies", "VERIFIED", "Node v20/v22 compatibility, npm packages installed without vulnerabilities."),
        ("Unit Test Suite", "VERIFIED", "14/14 unit tests passed (node:test) covering QR toggle, safe redirects, and limits."),
        ("Application Startup", "VERIFIED", "Next.js 16.2.6 dev server running cleanly on port 3001 with Turbopack."),
        ("Health Diagnostic Endpoint", "VERIFIED", "GET /api/health returned 200 OK with mongodbConfigured: true and jwtConfigured: true."),
        ("Admin Authentication Flow", "VERIFIED", "Created admin@ug.edu.gh via CLI; logged in and issued 24h JWT cookie."),
        ("Student Authentication Flow", "VERIFIED", "Created student STU001; verified student ID login and set-password flow."),
        ("Geofence Map & Sites", "VERIFIED", "Legon Campus (Lat 5.6502, Lon -0.1962, Radius 500m) created and rendered on Leaflet."),
        ("Course Timetables & QR", "VERIFIED", "Created CS201 & CS202; verified single-code dual-scan QR generation and enrollment."),
        ("UI Screenshot Evidence", "VERIFIED", "17 high-resolution real UI screenshots captured and embedded directly into Word guide."),
        ("Production Compilation", "VERIFIED", "Optimized production build generated (53/53 static/dynamic routes compiled successfully).")
    ]

    for i, row in enumerate(v_data, start=1):
        for j, val in enumerate(row):
            cell = v_table.cell(i, j)
            bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            r = cell.paragraphs[0].add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if j == 0:
                r.bold = True
                r.font.color.rgb = COLOR_NAVY
            elif j == 1:
                r.font.color.rgb = COLOR_GREEN
                r.bold = True
            else:
                r.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Concluding Signature Box
    add_callout(doc, [
        "This concludes the complete technical handover and user manual for UG Attend (ug_attend).",
        "The project is fully operational, verified, and ready for deployment at the University of Ghana."
    ], box_type="success", title="Verification Complete")

    # Save document
    doc.save(DOC_PATH)
    print(f"\n=======================================================")
    print(f"SUCCESS: Document successfully created at:")
    print(f"{DOC_PATH}")
    print(f"File Size: {os.path.getsize(DOC_PATH):,} bytes")
    print(f"=======================================================\n")

if __name__ == "__main__":
    build_document()
